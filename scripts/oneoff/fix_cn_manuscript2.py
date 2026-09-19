# -*- coding: utf-8 -*-
"""Second pass: regex-strip CI parentheticals and fix the abstract risk number."""
import re
from docx import Document

PATH = r"E:\class\Doctor\term4\Comp1\MedTrajectory_AI_Competition\docs\articles\Manuscript_cn_v1.docx"

# CI parentheticals: （95% CI ...） / （0.02--0.04） / （-0.039 至 -0.023）
CI_RE = re.compile(r'（(?:95% CI\s*)?[0-9.\-−–+]*\s*(?:--|–|至)\s*[0-9.\-−–+]*）')
CI_CI = re.compile(r'（95% CI [^）]*）')


def clean_text(text):
    t = CI_CI.sub('', text)
    t = CI_RE.sub('', t)
    # trailing whitespace from removals
    t = re.sub(r'\s+（', '（', t)
    t = re.sub(r'\s{2,}', ' ', t)
    # abstract fix
    t = t.replace('Fixed-horizon risk AUROC changed by only 0.000272',
                  'Fixed-horizon risk discrimination was non-inferior')
    return t


def process(paragraph):
    runs = paragraph.runs
    if not runs:
        return 0
    text = paragraph.text
    new_text = clean_text(text)
    if new_text != text:
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""
        return 1
    return 0


def main():
    doc = Document(PATH)
    n = 0
    for paragraph in doc.paragraphs:
        n += process(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    n += process(paragraph)
    doc.save(PATH)
    print("saved, paragraphs changed:", n)


if __name__ == "__main__":
    main()
