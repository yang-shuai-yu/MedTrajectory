from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MedTrajectory_TrackG_论文草稿_v1_20260821.docx"

BLUE = "2E5D7B"
LIGHT = "EAF0F4"
GRAY = "666666"
INK = "1F2933"


def set_font(run, latin="Arial", east="Microsoft YaHei", size=None, bold=None, color=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tbl_header = OxmlElement("w:tblHeader")
                tr_pr.append(tbl_header)
            tbl_header.set(qn("w:val"), "true")
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, caption, headers, rows, widths, font_size=8.3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(caption)
    set_font(r, size=9, bold=True, color=INK)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_font(r, size=font_size, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLUE)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, size=10.5, bold=True, color=INK)
        text = text[len(bold_lead):]
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, size=8.5, color=GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.header_distance = Inches(0.45)
section.footer_distance = Inches(0.45)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 11.5, 8, 4)):
    style = doc.styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("MedTrajectory manuscript draft | 2026-08-21")
set_font(r, size=8, color=GRAY)
add_page_number(section.footer.paragraphs[0])

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(18)
title.paragraph_format.space_after = Pt(8)
r = title.add_run("面向纵向电子健康记录的时间感知生成式轨迹建模：\n预注册锁定测试、死亡事件解耦与删失感知风险评估")
set_font(r, size=20, bold=True, color=BLUE)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Time-aware Generative Modeling of Longitudinal Electronic Health Records with Prespecified Locked Testing and Censor-aware Mortality Risk")
set_font(r, size=11, italic=True, color=GRAY)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(8)
r = meta.add_run("作者：[待补]  |  单位：[待补]  |  通讯作者：[待补]  |  稿件版本：v1, 2026-08-21")
set_font(r, size=9.5, color=GRAY)

add_heading(doc, "摘要", 1)
add_body(doc, "背景：纵向电子健康记录（EHR）生成模型需要同时保持事件内容、事件时间和轨迹长度的真实性，但将死亡作为普通词表事件可能使死亡频率与轨迹终止机制相互混淆。方法：我们构建了时间感知的自回归轨迹模型，并将预先选定的 A2 主模型与 A0 架构控制、ETHOS-Matched 和 Foresight-Matched 进行同协议比较。在 1,669 名患者的 one-shot locked-test 中，每位患者进行 20 次 rollout，并以患者级配对 bootstrap 评估诊断集合重合、下一事件检索、首事件时间和事件数量。随后仅在 validation 内开展去 death-token 消融及独立离散时间 death-hazard head 实验，后者使用正式随访终点处理右删失，并在固定内部二分中训练、校准和评估。结果：相对 A0，A2 的 Diagnosis Jaccard 提高 0.01385（95% CI 0.01065–0.01706），Hit@10 提高 0.03215（0.02077–0.04394），首事件时间 MAE 降低 5.01 天（−6.57 至 −3.44），事件数量 MAE 降低 0.899（−1.007 至 −0.787）。直接去除 death token 会使 A2 的 Jaccard 从 0.0834 降至 0.0530，并使 event-count MAE 从 12.9335 升至 24.2815，说明死亡事件同时承担了不完整的轨迹终止作用。独立删失感知 hazard head 在相同 187 名 10 年可评估患者上将 A2 的 Brier score 改善 0.0970（95% CI 0.0367–0.1552）。结论：A2 对纵向轨迹保真度提供稳定且可重复的结构性增益；死亡风险应从普通事件生成中解耦，并通过独立的删失感知风险头与预先规定的终止门控建模。")

p = doc.add_paragraph()
r = p.add_run("关键词：")
set_font(r, size=10, bold=True, color=INK)
r = p.add_run("电子健康记录；患者轨迹生成；生成式 Transformer；相对时间编码；生存分析；右删失；死亡校准")
set_font(r, size=10, color=INK)

add_heading(doc, "1 引言", 1)
add_body(doc, "纵向 EHR 包含诊断、操作、肿瘤登记、死亡和静态协变量等异质事件。它既是不规则采样的时间序列，也是具有临床语义约束的事件序列。BEHRT 与 Med-BERT 表明 Transformer 可以从结构化 EHR 学习可迁移的患者表示并支持疾病预测[4,5]；Foresight、ETHOS 和 Delphi-2M 则进一步把临床记录表示为可生成的患者时间线或健康轨迹[1–3]。这些工作共同说明，大规模序列预训练能够承载多任务风险预测与未来轨迹模拟。")
add_body(doc, "然而，轨迹生成与死亡风险并非同一统计任务。生成模型通常把死亡表示为一个离散 token；患者层面的“死亡概率”继而被估计为多次 rollout 中出现 death token 的频率。这一概率同时受事件 softmax、采样温度、生成长度和终止规则影响，未显式处理右删失，也不必然具有绝对风险校准。相反，DeepHit、Dynamic-DeepHit、SurvTRACE 及后续 time-to-event foundation models 将删失和时间窗作为风险建模的核心[6–10]。因此，把死亡 token 生成频率直接解释为临床死亡风险可能造成任务错配。")
add_body(doc, "本研究提出并检验一个分层框架：首先用生成式主干学习未来临床事件和相对时间结构；其次把死亡从普通事件概率中诊断性地移除，以检验轨迹收益是否由 death token 驱动；最后使用独立的 censor-aware hazard head 建模 1、5 和 10 年死亡风险。我们遵循 validation-only 开发和 one-shot locked-test 原则，避免使用 locked-test 结果调参。研究重点不是宣称单一模型在所有指标上获胜，而是区分轨迹保真度、时间拟合、死亡判别与概率校准。")

add_heading(doc, "2 相关工作与研究定位", 1)
lit_rows = [
    ("Delphi-2M", "生成式疾病自然史；年龄条件轨迹", "最直接的生成范式参照", "PMC12657216"),
    ("Foresight", "生成式患者时间线；EHR 文本/概念", "matched baseline；非逐层复刻", "PMC11220626"),
    ("ETHOS", "零样本健康轨迹与结局预测", "matched baseline；死亡频率较稳", "PMC11412988"),
    ("BEHRT", "双向 Transformer EHR 表征", "项目内 no-leak encoder baseline", "PMC7189231"),
    ("Med-BERT", "结构化 EHR 预训练与疾病预测", "风险分类路线参照", "PMC8137882"),
    ("DeepHit / Dynamic-DeepHit", "离散时间、竞争风险、动态预测", "独立 hazard head 的统计学邻近路线", "非 PMC 核心会议/IEEE"),
    ("SurvTRACE", "Transformer survival / competing risks", "支持 censor-aware 多时间窗建模", "PMC11185454（应用验证）"),
    ("EMR adaptive risk foundation model", "预训练表示 + 多风险预测", "支持生成表征与风险头解耦", "PMC12482913"),
]
add_table(doc, "表 1. 与本研究思路最接近的文献路线", ["工作", "核心方法", "与本研究关系", "PMC 可核验性"], lit_rows, [1600, 2700, 3000, 2060], 7.6)
add_body(doc, "文献定位结论：现有工作已经分别覆盖“生成式健康轨迹”和“删失感知 time-to-event 风险”两端，但将 death 从普通事件词表中诊断性解耦、进行患者级 raw/no-death 配对消融，再用固定 validation 内部二分训练独立 hazard head，并把该结果与预注册 locked-test 轨迹指标共同报告的组合证据仍较少见。本文的贡献主要来自任务分解、严格比较协议和实证闭环，而不是声称首次提出 survival head。")

add_heading(doc, "3 方法", 1)
add_heading(doc, "3.1 研究设计与数据划分", 2)
add_body(doc, "研究包含 Track R（输入与结构消融）和 Track G（多步生成与死亡建模）两个阶段。A2 在 Track R v2.2 中依据 validation 结果预先选定并冻结，随后作为 Track G 的 primary candidate。locked-test 包含 1,669 名患者，其中 1,541 名患者具有有效 waiting-time 评价数据。四个模型、三个随机种子（42、43、44）共形成 12 个 locked-test 运行，每位患者进行 20 次 rollout。测试读取是 one-shot；任何后续校准、阈值、采样或 gate 设计均仅允许使用 validation。")
add_body(doc, "[待补：数据来源、纳入排除标准、训练/validation/test 患者数、观察窗口、预测起点、事件编码规范、伦理审批编号、知情同意豁免与数据访问条件。不得在稿件中写入可识别参与者信息。]")

add_heading(doc, "3.2 事件表示与模型", 2)
add_body(doc, "动态输入统一包含 diagnosis、procedure、cancer 和 death 等事件，静态信息通过前缀 token 注入。主干采用 decoder-only Transformer，并联合预测下一事件及事件间隔。A0 保留绝对年龄 Sin/Cos 表示，作为架构控制；A2 在保留绝对年龄的同时，对动态事件间 attention 加入 additive continuous-age relative RoPE，使注意力分数直接感知真实年龄差。该设计避免以相对位置完全替换绝对年龄而损失生命周期信息。")
add_body(doc, "ETHOS-Matched 与 Foresight-Matched 在相同 Track G 数据目录、患者划分、静态前缀、dynamic BOS、动态事件类型、词表、mask 和训练预算下实现。它们是协议受控的 matched implementations，不是对原论文预训练规模、所有特征工程和逐层架构的完全复现。A0/A2 使用历史 Track R v2.2 checkpoint，而 matched baselines 在 Track G 中重新训练；因此外部模型比较用于任务级基准，不作为纯架构因果估计。")

add_heading(doc, "3.3 轨迹生成与评价", 2)
add_body(doc, "轨迹保真度由四个互补指标定义：Diagnosis Jaccard 衡量预测与真实诊断集合的重合；Hit@10 衡量真实下一事件是否进入前 10 候选；First-event time MAE 衡量首个未来事件的时间定位误差；Event-count MAE 衡量生成事件数量与真实数量的偏差。Waiting-time NLL 作为辅助时间指标单列，死亡 Brier 仅描述 rollout-derived death probability，不并入轨迹保真度主结论。")
add_body(doc, "模型差值先在同一患者内跨三个 seed 聚合，再进行患者级配对 bootstrap。locked-test 主分析使用 10,000 次 bootstrap；validation follow-up 使用 2,000 次 bootstrap。所有置信区间均为 95%。")

add_heading(doc, "3.4 死亡事件消融与独立风险头", 2)
add_body(doc, "在 validation 的 1,000 名患者上，我们固定原 sampler 和每患者 20 次 rollout，将 death 同时从普通事件候选和普通事件时间速率中移除。该实验不重训主模型，目的在于检验 A2 相对 A0 的收益是否依赖 death token，并评估缺失终止机制对轨迹长度的影响。")
add_body(doc, "独立 death-hazard head 冻结四个模型的主干表示，以离散时间 hazard 预测 1、5 和 10 年死亡风险。validation 使用固定 hash 二分：498 人进入 outer fit，502 人进入 final eval；outer fit 再固定分为 262 人 head-train 和 236 人 calibration-fit。损失显式处理右删失，随访终点来自正式的 val_followup_end_age_days.npy。共享 Platt 变换仅在 calibration-fit 拟合，最终指标仅在 eval 计算。")

add_heading(doc, "4 结果", 1)
add_heading(doc, "4.1 已完成的比较体系", 2)
comparison_rows = [
    ("输入信息消融", "Disease-only、+procedure、+cancer/death/static", "已完成", "多类型事件是主要增益来源"),
    ("结构控制", "A2 vs A0，三 seed", "已完成", "检验 additive relative RoPE"),
    ("外部路线 matched 对照", "A2 vs ETHOS-Matched / Foresight-Matched", "已完成", "同协议，不是原论文完整复现"),
    ("跨拆分复现", "validation vs one-shot locked-test", "已完成", "主轨迹指标方向一致"),
    ("死亡 token 消融", "raw vs no-death，患者级配对", "已完成", "直接删除破坏终止机制"),
    ("死亡风险解耦", "rollout probability vs independent hazard head", "已完成", "同一 censor-valid 患者比较"),
    ("原论文级公平复现", "原始预训练语料/规模/全部特征", "未完成", "不能宣称全面击败原 ETHOS/Foresight"),
    ("hazard gate 回接 rollout", "独立风险控制生成终止", "待开展", "须 validation-only 预注册"),
]
add_table(doc, "表 2. 本项目已经完成与尚未完成的比较", ["层次", "对比", "状态", "解释边界"], comparison_rows, [1500, 2900, 1000, 3960], 7.6)

add_heading(doc, "4.2 Track R 输入与结构消融", 2)
track_r_rows = [
    ("Disease-only", "diagnosis", "0.2935", "0.7274"),
    ("Exp1", "diagnosis + procedure", "0.4832", "0.7949"),
    ("Exp2", "diagnosis + procedure + cancer + death", "0.4891", "0.7972"),
]
add_table(doc, "表 3. 早期输入消融", ["模型", "动态输入", "Full Top10", "Selected-disease mean AUROC"], track_r_rows, [1600, 3900, 1700, 2160], 8.1)
add_body(doc, "正式 P0–P4 矩阵中，完整多类型事件输入 P2 相对 P0 的 LM AUROC 增量为 0.127029（95% CI 0.123193–0.130864）。Track R v2.1 同时显示，静态前缀相对无静态信息模型提高宏 AUROC 0.026818，而旧 CARoPE 与静态前缀存在负 interaction（−0.061026，95% CI −0.068845 至 −0.052912）。这些结果促使 v2.2 采用“绝对年龄 + additive relative RoPE”，而非用相对时间完全替换绝对年龄。")

add_heading(doc, "4.3 One-shot locked-test 轨迹结果", 2)
locked_rows = [
    ("A0", "0.05721 ± 0.00198", "0.37547 ± 0.00840", "302.46 ± 0.73", "13.691 ± 0.765", "0.42110"),
    ("A2", "0.07107 ± 0.00124", "0.40763 ± 0.00540", "297.44 ± 1.70", "12.792 ± 1.035", "0.49986"),
    ("ETHOS-Matched", "0.04879 ± 0.00408", "0.38107 ± 0.00864", "302.51 ± 0.49", "20.298 ± 0.698", "0.12285"),
    ("Foresight-Matched", "0.03228 ± 0.00208", "0.37328 ± 0.00452", "302.18 ± 3.55", "21.903 ± 0.314", "0.08684"),
]
add_table(doc, "表 4. Track G locked-test 三 seed 均值（N=1,669）", ["模型", "Jaccard", "Hit@10", "First-event MAE, d", "Count MAE", "Death Brier"], locked_rows, [1500, 1800, 1550, 1800, 1500, 1210], 7.2)
delta_rows = [
    ("A2 − A0", "+0.01385 [0.01065, 0.01706]", "+0.03215 [0.02077, 0.04394]", "−5.01 [−6.57, −3.44]", "−0.899 [−1.007, −0.787]"),
    ("A2 − ETHOS-Matched", "+0.02227 [0.01778, 0.02686]", "+0.02656 [0.01398, 0.03934]", "−5.07 [−7.39, −2.73]", "−7.506 [−7.736, −7.277]"),
    ("A2 − Foresight-Matched", "+0.03878 [0.03204, 0.04560]", "+0.03435 [0.02017, 0.04873]", "−4.73 [−7.11, −2.37]", "−9.111 [−9.389, −8.832]"),
]
add_table(doc, "表 5. A2 的患者级配对效应量及 95% CI", ["比较", "Jaccard", "Hit@10", "First-event MAE, d", "Count MAE"], delta_rows, [1650, 2250, 2200, 1700, 1560], 7.2)
add_body(doc, "A2 相对 A0 的四个主要轨迹指标置信区间均排除 0，且方向与 validation 一致。因此，我们将 A2 的收益表述为“纵向轨迹保真度的稳定改善”。Foresight-Matched 在 death Brier 和 waiting-time NLL 上更优，说明不同模型在事件内容保真度、时间似然和死亡概率尺度之间存在权衡。")

add_heading(doc, "4.4 死亡概率失准与 no-death-token 消融", 2)
add_body(doc, "locked-test 的真实死亡率为 0.122，而 A0、A2、ETHOS-Matched 和 Foresight-Matched 的平均 rollout death probability 分别约为 0.653、0.729、0.238 和 0.079。A0/A2 具有死亡排序信号，但概率尺度严重偏高；完整 validation 上 A0/A2 的 AUROC 分别为 0.8446 和 0.8298，而 Brier skill 分别为 −1.7962 和 −2.4327。相对地，Foresight-Matched 的 AUROC 为 0.8207、Brier skill 为 0.2749，表明其优势主要体现在生成频率和校准，而不是所有轨迹任务。")
no_death_rows = [
    ("A0", "Jaccard", "0.0697", "0.0378", "−0.0319 [−0.0409, −0.0242]"),
    ("A2", "Jaccard", "0.0834", "0.0530", "−0.0305 [−0.0390, −0.0227]"),
    ("A0", "Hit@10", "0.4057", "0.4117", "+0.0060 [+0.0030, +0.0090]"),
    ("A2", "Hit@10", "0.4434", "0.4478", "+0.0043 [+0.0020, +0.0070]"),
    ("A0", "Count MAE", "14.0672", "24.3305", "+10.2633 [+9.8678, +10.6637]"),
    ("A2", "Count MAE", "12.9335", "24.2815", "+11.3481 [+10.9566, +11.7636]"),
]
add_table(doc, "表 6. Validation-only 去 death-token 配对消融", ["模型", "指标", "Raw", "No-death", "Delta [95% CI]"], no_death_rows, [1200, 1500, 1200, 1400, 4060], 7.8)
add_body(doc, "去除 death token 后，下一事件检索略有改善，但 Diagnosis Jaccard 和事件数量误差显著恶化。原因是生成器失去终止机制，更容易持续生成至 max_new_tokens。重要的是，no-death 条件下 A2 相对 A0 的 Jaccard 优势仍为 0.0151（95% CI 0.0117–0.0189），Hit@10 优势为 0.0360（0.0204–0.0511），说明 A2 的结构收益并非仅由 death token 生成驱动。")

add_heading(doc, "4.5 独立删失感知 death-hazard head", 2)
hazard_rows = [
    ("A0", "1y", "495 / 37", "0.9313", "0.5152", "0.0541", "0.0585"),
    ("A2", "1y", "495 / 37", "0.9225", "0.4884", "0.0559", "0.0644"),
    ("ETHOS-Matched", "1y", "495 / 37", "0.9278", "0.5597", "0.0546", "0.0672"),
    ("Foresight-Matched", "1y", "495 / 37", "0.9231", "0.5049", "0.0552", "0.0561"),
    ("A0", "5y", "366 / 78", "0.8272", "0.6000", "0.1234", "0.0427"),
    ("A2", "5y", "366 / 78", "0.8406", "0.6857", "0.1088", "0.0427"),
    ("ETHOS-Matched", "5y", "366 / 78", "0.8346", "0.6948", "0.1083", "0.0707"),
    ("Foresight-Matched", "5y", "366 / 78", "0.7934", "0.5992", "0.1267", "0.0409"),
    ("A0", "10y", "187 / 88", "0.8910", "0.9103", "0.1240", "0.0869"),
    ("A2", "10y", "187 / 88", "0.8379", "0.8493", "0.1589", "0.0930"),
    ("ETHOS-Matched", "10y", "187 / 88", "0.8743", "0.8931", "0.1367", "0.0672"),
    ("Foresight-Matched", "10y", "187 / 88", "0.8565", "0.8691", "0.1488", "0.0755"),
]
add_table(doc, "表 7. Independent hazard head 的 final-eval 结果", ["模型", "窗", "N / deaths", "AUROC", "AUPRC", "Brier", "ECE"], hazard_rows, [1600, 650, 1300, 1250, 1250, 1250, 1060], 7.4)
add_body(doc, "在相同的 187 名 10 年 censor-valid 患者上，独立 head 相对原 rollout 的 Brier 差值（head − rollout）为：A0 −0.0920（95% CI −0.1481 至 −0.0379）、A2 −0.0970（−0.1552 至 −0.0367）、ETHOS-Matched −0.0235（−0.0543 至 0.0078）、Foresight-Matched −0.1045（−0.1522 至 −0.0580）。这表明独立 hazard head 明确改善 A0、A2 和 Foresight-Matched 的概率误差；ETHOS-Matched 方向一致但置信区间跨 0。")

add_heading(doc, "5 讨论", 1)
add_heading(doc, "5.1 A2 的有效性应如何表述", 2)
add_body(doc, "A2 是 Track G 之前预先选定的主模型，而不是在 locked-test 后挑选的最好模型。它相对 A0 在四个互补轨迹指标上均有统计显著改善，且 no-death 条件下仍保留优势。因此，现有证据足以支持“additive continuous-age relative attention 改善纵向轨迹保真度”。但是，A2 并未在 death calibration、waiting-time NLL 或所有 horizon risk 上占优，不能表述为全面优于所有基线。")

add_heading(doc, "5.2 为什么 Foresight 的死亡结果更好", 2)
add_body(doc, "Foresight-Matched 的低 death Brier 主要来自更接近真实发生率的生成频率和较好的概率尺度。A0/A2 的 AUROC 与其相当或更高，说明主干中存在患者风险排序信号；问题在于这个信号通过普通事件采样、多步累积和终止规则被放大。换言之，Foresight 的优势更接近 calibration/rollout behavior 优势，而不是对所有未来临床事件的全面建模优势。这也解释了为何 A2 可同时具有更高轨迹 Jaccard 与更差 death Brier。")

add_heading(doc, "5.3 与现有文献相比的贡献", 2)
add_body(doc, "与 Delphi-2M、Foresight 和 ETHOS 相比，本研究强调多类型事件输入、绝对年龄与相对时间的联合表示，以及患者级 paired trajectory metrics。与 DeepHit、Dynamic-DeepHit 和 SurvTRACE 相比，本研究并未提出新的生存损失，而是把删失感知风险头置于生成式轨迹模型之后，作为对 rollout death probability 的校正与任务解耦。最有说服力的新意是：同一研究协议内同时展示生成轨迹的结构收益、death token 的终止作用、直接删除的失败，以及独立风险头对概率误差的改善。")

add_heading(doc, "5.4 局限性", 2)
for item in [
    "ETHOS-Matched 和 Foresight-Matched 是同协议重实现，未复制原论文的全部预训练语料、模型规模和特征工程，因此不能宣称本文模型全面超过原始发表模型。",
    "A0/A2 使用历史 checkpoint，而 matched baselines 在 Track G 中重新训练；虽然训练预算和评估协议受控，但并非完全相同初始化下的纯结构实验。",
    "独立 hazard head 仅在 validation 内开发和评估，尚无第二个 untouched test 对其泛化和临床校准进行确认。",
    "10 年风险仅有 187 名 censor-valid eval 患者、88 例死亡，长期窗估计的不确定性较大。",
    "当前尚未将 hazard head 回接生成器作为 termination gate；因此还不能证明解耦设计能同时改善死亡校准与完整轨迹保真度。",
    "数据来源、编码质量、医疗利用差异和人群迁移可能影响外部效度，且生成结果不应直接用于临床决策。",
]:
    add_bullet(doc, item)

add_heading(doc, "6 结论", 1)
add_body(doc, "在预先规定的 one-shot locked-test 中，A2 相对 A0 稳定改善诊断集合重合、下一事件检索、首事件时间定位与事件数量准确性。该优势在移除 death token 后仍存在，排除了“收益完全由死亡生成驱动”的解释。另一方面，直接删除 death token 会破坏轨迹终止并显著恶化整体保真度。独立、删失感知的 hazard head 能降低 rollout-derived death probability 的误差，支持将轨迹生成与死亡风险拆分为互补任务。下一步应仅在 validation 中预注册 hazard gate 形式、时间窗和停止规则，再评估其对完整 rollout 的影响。")

add_heading(doc, "7 预注册的下一步分析", 1)
for item in [
    "固定 1/5/10 年 discrete-time hazard 输出，不根据 locked-test 选择阈值。",
    "定义 hazard-to-stop 转换、竞争事件与随访终点处理，并在代码和配置中冻结。",
    "在 validation 上进行相同患者、相同 sampler 的 raw vs gated paired trajectory comparison。",
    "同时报告 Jaccard、Hit@10、首事件时间、事件数量、death Brier/ECE、AUROC/AUPRC，避免用单一指标选择模型。",
    "若未来获得新的 untouched cohort，再对 gate 做一次独立确认；当前 locked-test 不重跑。",
]:
    add_bullet(doc, item)

add_heading(doc, "声明", 1)
add_body(doc, "伦理审批：[待补]。数据可用性：[待补；需符合数据提供方协议]。代码可用性：[待补仓库/归档 DOI]。利益冲突：[待补]。资助信息：[待补]。作者贡献：[待补，建议按 CRediT taxonomy]。本文所有患者级分析仅报告聚合统计，不记录或披露参与者标识。")

add_heading(doc, "参考文献", 1)
refs = [
    "1. Shmatko A, Jung AW, Gaurav K, et al. Learning the natural history of human disease with generative transformers. Nature. 2025. doi:10.1038/s41586-025-09529-3. PMCID: PMC12657216.",
    "2. Kraljevic Z, Bean D, Shek A, et al. Foresight—a generative pretrained transformer for modelling of patient timelines using electronic health records: a retrospective modelling study. Lancet Digital Health. 2024;6:e281–e290. doi:10.1016/S2589-7500(24)00025-6. PMCID: PMC11220626.",
    "3. Renc P, Jia Y, Samir A, et al. Zero shot health trajectory prediction using transformer. npj Digital Medicine. 2024;7. doi:10.1038/s41746-024-01235-0. PMCID: PMC11412988.",
    "4. Li Y, Rao S, Solares JRA, et al. BEHRT: Transformer for Electronic Health Records. Scientific Reports. 2020;10:7155. doi:10.1038/s41598-020-62922-y. PMCID: PMC7189231.",
    "5. Rasmy L, Xiang Y, Xie Z, Tao C, Zhi D. Med-BERT: pretrained contextualized embeddings on large-scale structured electronic health records for disease prediction. npj Digital Medicine. 2021;4:86. doi:10.1038/s41746-021-00455-y. PMCID: PMC8137882.",
    "6. Lee C, Zame WR, Yoon J, van der Schaar M. DeepHit: A Deep Learning Approach to Survival Analysis With Competing Risks. AAAI. 2018;32. doi:10.1609/aaai.v32i1.11842.",
    "7. Lee C, Yoon J, van der Schaar M. Dynamic-DeepHit: A Deep Learning Approach for Dynamic Survival Analysis With Competing Risks Based on Longitudinal Data. IEEE Trans Biomed Eng. 2020;67:122–133. doi:10.1109/TBME.2019.2909027.",
    "8. Wang Z, Sun J. SurvTRACE: Transformers for Survival Analysis with Competing Events. ACM BCB. 2022. doi:10.1145/3535508.3545521.",
    "9. Ishihara S, et al. The potential of the transformer-based survival analysis model, SurvTrace, for predicting recurrent cardiovascular events. PLoS ONE. 2024;19:e0304423. doi:10.1371/journal.pone.0304423. PMCID: PMC11185454.",
    "10. Renc P, Grzeszczyk MK, Oufattole N, et al. Foundation model of electronic medical records for adaptive risk estimation. GigaScience. 2025;14:giaf107. doi:10.1093/gigascience/giaf107. PMCID: PMC12482913.",
    "11. Su J, Lu Y, Pan S, Murtadha A, Wen B, Liu Y. RoFormer: Enhanced Transformer with Rotary Position Embedding. Neurocomputing. 2024;568:127063. doi:10.1016/j.neucom.2023.127063.",
    "12. Vaswani A, Shazeer N, Parmar N, et al. Attention Is All You Need. Advances in Neural Information Processing Systems. 2017;30.",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(ref)
    set_font(r, size=8.8, color=INK)

add_heading(doc, "附录 A：论文图表建议", 1)
for item in [
    "Figure 1：研究设计图，展示 Track R 选择 A2、Track G locked-test、validation-only death follow-up 三个阶段及信息隔离。",
    "Figure 2：A2−A0、A2−ETHOS-Matched、A2−Foresight-Matched 的四项 paired effect forest plot。",
    "Figure 3：四模型 rollout death probability 的 reliability curves，与独立 hazard head 的 1/5/10 年 ECE/Brier 对照。",
    "Extended Data Figure 1：raw/no-death 的患者级 delta 分布与轨迹长度分布。",
    "Extended Data Table 1：逐 seed 全部指标；Extended Data Table 2：患者数、删失有效数、协议 hash 与 QA；Extended Data Table 3：death head 的 horizon-specific calibration。",
]:
    add_bullet(doc, item)

doc.core_properties.title = "MedTrajectory Track G manuscript draft"
doc.core_properties.subject = "Longitudinal EHR trajectory generation and censor-aware mortality risk"
doc.core_properties.author = "[Authors pending]"
doc.core_properties.keywords = "EHR, trajectory generation, survival analysis, mortality calibration"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
