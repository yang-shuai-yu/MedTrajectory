# -*- coding: utf-8 -*-
"""Append the locked-test version of Supplementary Table S7 to both Supplementary files."""
from docx import Document


def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    return t


HEADER = ["Model", "Macro AUROC ↑", "Macro AUPRC ↑", "Brier ↓", "ECE ↓"]
ROWS = [
    ["MedTrajectory-Relative (A2)", "0.7499", "0.2611", "0.0313", "0.0322"],
    ["MedTrajectory-Absolute (A0)", "0.7452", "0.2637", "0.0291", "0.0304"],
    ["Cox-R", "0.7265", "0.2061", "0.0416", "0.0366"],
    ["MDRMF-Clinical-R", "0.7205", "0.2133", "0.0280", "0.0244"],
    ["Logistic-R", "0.7159", "0.2044", "0.2794", "0.4613"],
    ["Med-BERT-Paper", "0.7045", "0.2010", "0.0353", "0.0380"],
    ["Med-BERT-Matched-S", "0.6650", "0.1699", "0.0506", "0.0535"],
]


def append_cn(path):
    doc = Document(path)
    doc.add_page_break()
    doc.add_heading("附：Supplementary Table S7（locked test 版）——固定期限风险基准", level=1)
    doc.add_paragraph(
        "在锁定风险测试集（Track R test，N=6,099）上重新评估的完整风险基准。"
        "MedTrajectory 为 3 个随机种子（42/43/44）均值；基线（Cox/Logistic/MDRMF/Med-BERT）为单 seed 42 的冻结 checkpoint，"
        "直接在 test 特征上预测。本表取代原 validation 版 Table S7。"
    )
    add_table(doc, HEADER, ROWS)
    doc.add_paragraph(
        "注：(1) 原 validation 版 Table S7 中的 no-static 与 legacy CARoPE（A1）消融行未在 test 上重跑，"
        "仍以 validation 数值为准（见原表），本表聚焦 MedTrajectory 与经典基线的核心对比；"
        "(2) Logistic-R 的 checkpoint 以 scikit-learn 1.8.0 拟合，本次用 1.9.0 加载预测（系数固定，预测不受影响）；"
        "(3) 口径为 sex×disease×horizon×age 宏平均。"
    )
    doc.save(path)
    print("saved cn:", path)


def append_en(path):
    doc = Document(path)
    doc.add_page_break()
    doc.add_heading("Appendix: Supplementary Table S7 (locked-test version) — fixed-horizon risk benchmark", level=1)
    doc.add_paragraph(
        "Complete risk benchmark re-evaluated on the locked risk test set (Track R test, N = 6,099). "
        "MedTrajectory values are means over three random seeds (42/43/44); the baselines "
        "(Cox/Logistic/MDRMF/Med-BERT) use single-seed-42 frozen checkpoints predicting on test features. "
        "This table supersedes the original validation-based Table S7."
    )
    add_table(doc, HEADER, ROWS)
    doc.add_paragraph(
        "Notes: (1) the no-static and legacy CARoPE (A1) ablation rows from the original validation Table S7 were not "
        "re-run on test and remain validation values; this table focuses on the core MedTrajectory-vs-classical-baseline "
        "comparison. (2) The Logistic-R checkpoint was fitted with scikit-learn 1.8.0 and loaded with 1.9.0 for prediction "
        "(coefficients are fixed, so predictions are unaffected). (3) Aggregation is macro over sex × disease × horizon × age."
    )
    doc.save(path)
    print("saved en:", path)


if __name__ == "__main__":
    append_cn(r"E:\class\Doctor\term4\Comp1\MedTrajectory_AI_Competition\docs\articles\Supplementary_cn_v1.docx")
    append_en(r"E:\class\Doctor\term4\Comp1\MedTrajectory_AI_Competition\docs\articles\Supplementary_en_v1.docx")
