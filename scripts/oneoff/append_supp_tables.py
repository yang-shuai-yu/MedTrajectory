# -*- coding: utf-8 -*-
"""Append the resolved tables/params to the end of Supplementary_cn_v1.docx
without touching existing content. Run with local python (python-docx 1.2.0).
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = r"E:\class\Doctor\term4\Comp1\MedTrajectory_AI_Competition\docs\articles\Supplementary_cn_v1.docx"


def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    return t


def main():
    doc = Document(PATH)
    doc.add_page_break()

    doc.add_heading("附：待补数据与参数（本次从服务器结果补齐）", level=1)

    # ---- A. P0-P4 matrix ----
    doc.add_heading("A. P0–P4 完整输入消融矩阵（逐 seed）", level=2)
    doc.add_paragraph(
        "配置定义：P0 = diagnosis_death（诊断+死亡，Delphi 兼容基线，无 RoPE）；"
        "P1 = diagnosis_death + continuous-age RoPE；P2 = multitype（诊断+操作+肿瘤+死亡，无 RoPE）；"
        "P3 = multitype + RoPE（主模型）；P4 = reduced_multitype + RoPE（队列规模对照）。"
        "数据源：results/paper_protocol_v1/matrix_seeds42-44_my3_20260806_0252/。"
    )
    doc.add_heading("Table A1. LM AUROC（lm_mean_auc）逐 seed 与均值", level=3)
    add_table(doc,
              ["配置", "seed42", "seed43", "seed44", "均值"],
              [["P0", "0.683329", "0.682850", "0.683578", "0.683252"],
               ["P1", "0.681108", "0.680659", "0.682650", "0.681472"],
               ["P2", "0.808919", "0.811511", "0.810414", "0.810281"],
               ["P3", "0.809723", "0.811335", "0.810685", "0.810581"],
               ["P4", "0.809434", "0.809046", "0.809899", "0.809460"]])
    doc.add_paragraph(
        "关键效应：P2 − P0 = +0.127029（95% CI 0.123193–0.130864），与正文 §5.1 一致。"
    )
    doc.add_heading("Table A2. Horizon 风险宏 AUROC（horizon_risk_mean_auc）逐 seed 与均值", level=3)
    add_table(doc,
              ["配置", "seed42", "seed43", "seed44", "均值"],
              [["P0", "0.726366", "0.710024", "0.722343", "0.719578"],
               ["P1", "0.719345", "0.721297", "0.709984", "0.716875"],
               ["P2", "0.739660", "0.725068", "0.729119", "0.731282"],
               ["P3", "0.726463", "0.725191", "0.748698", "0.733451"],
               ["P4", "0.734302", "0.745536", "0.732584", "0.737474"]])

    # ---- B. Capacity scaling ----
    doc.add_heading("B. 容量扩展（A1-S / A1-M / A1-L）精确指标", level=2)
    doc.add_paragraph(
        "数据源：results/paper_protocol_v2/car_rope_capacity_comparison_20260812/"
        "（macro AUROC over sex×disease×horizon×age cells）。"
    )
    add_table(doc,
              ["模型（参数）", "val 宏 AUROC", "test 宏 AUROC"],
              [["A1-S（~1.39M）", "0.717275", "0.735256"],
               ["A1-M（3.67M）", "0.723697", "0.715972"],
               ["A1-L（4.03M）", "0.706786", "0.694866"]])
    doc.add_paragraph(
        "test 上 A1-M / A1-L 均低于 A1-S（Δ −0.019284 / −0.040389，Holm p 0.011988 / 0.005994），"
        "支持「不增加容量」的结论。"
    )

    # ---- C. Training config ----
    doc.add_heading("C. 训练与实现配置（补齐 Supplementary Table S15 的 Not recorded）", level=2)
    add_table(doc,
              ["Item", "Configuration"],
              [["Optimizer", "AdamW（β1=0.9, β2=0.99，CUDA fused）"],
               ["Batch size", "128"],
               ["Learning rate", "3e-4 → min 3e-5，cosine，warmup 500 iters"],
               ["Weight decay", "0.1"],
               ["Gradient clipping", "1.0（max norm）"],
               ["Dropout", "0.1"],
               ["Context length", "block_size 133（静态前缀 4 + BOS 1 + 动态 128）"],
               ["Training budget", "pretraining 100,000 iters；risk head 10,000 iters"],
               ["Random seeds", "42 / 43 / 44"],
               ["Numerical precision", "fp32（无 mixed precision）"],
               ["Hardware", "CUDA（CUDA_VISIBLE_DEVICES=0；火山云 RTX 4090 D）"]])

    # ---- D. Qwen ARI correction ----
    doc.add_heading("D. Qwen Phecode ARI 数值修正", level=2)
    doc.add_paragraph(
        "技术文档 Table 16 / Supplementary S13 中 Qwen 的 Phecode ARI 0.3121 为笔误。"
        "以 outputs/icd_embedding_compare/comparison.json 为准："
    )
    add_table(doc,
              ["指标（Qwen, 64d）", "正确值", "原表笔误值"],
              [["Phecode ARI", "0.3211", "0.3121"],
               ["Phecode NMI", "0.8636", "0.8642"]])

    # ---- E. Figure references ----
    doc.add_heading("E. Supplementary Figure 引用", level=2)
    add_table(doc,
              ["图", "文件路径"],
              [["Figure S1（死亡校准）", "docs/figures/figS1_death_calibration.png"],
               ["Figure S2（ICD 表示 UMAP）", "docs/figures/figS2_icd_umap.png"]])

    doc.save(PATH)
    print("saved:", PATH)
    print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))


if __name__ == "__main__":
    main()
